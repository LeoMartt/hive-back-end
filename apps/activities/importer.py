import csv
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from io import StringIO
from pathlib import Path
from typing import Any

from django.db import transaction
from rest_framework.exceptions import ValidationError as DRFValidationError

from apps.accounts.models import Usuario
from apps.projects.models import Membership, NoHierarquia, Papel, Project

from .models import Activity, ActivityPredecessor
from .services import obter_ou_criar_modulo, obter_ou_criar_processo, salvar_activity


HEADER_MAP = {
    "modulo": "module",
    "processo": "process",
    "atividade": "name",
    "atividadenome": "name",
    "nomedaatividade": "name",
    "nome": "name",
    "testeremail": "tester_email",
    "tester": "tester_email",
    "devemail": "developer_email",
    "desenvolvedoremail": "developer_email",
    "dev": "developer_email",
    "desenvolvedor": "developer_email",
    "datainicioplanejado": "planned_start",
    "datainicioplanejada": "planned_start",
    "inicioplanejado": "planned_start",
    "datainicio": "planned_start",
    "datafinalplanejada": "planned_end",
    "datafinalplanejado": "planned_end",
    "datafimplanejada": "planned_end",
    "conclusaoplanejada": "planned_end",
    "idlistasequencialtemporario": "temporary_id",
    "idlistasequencial": "temporary_id",
    "idtemporario": "temporary_id",
    "id": "temporary_id",
    "predecessores": "predecessors",
    "sistema": "system",
    "area": "area",
    "transacao": "transaction",
    "resultadoesperado": "expected_result",
    "observacoes": "notes",
    "observacao": "notes",
}

REQUIRED_FIELDS = [
    "module",
    "name",
    "tester_email",
    "developer_email",
    "planned_start",
    "planned_end",
    "temporary_id",
]


@dataclass(frozen=True)
class ImportedActivityRow:
    line: int
    temporary_id: int
    node: NoHierarquia
    name: str
    tester: Usuario
    developer: Usuario
    planned_start: date
    planned_end: date
    predecessor_temporary_ids: list[int]
    area: str
    system: str
    transaction: str
    expected_result: str
    notes: str


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_date(value: Any, line: int, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _clean(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise DRFValidationError({"rows": [f"Linha {line}: {field} inválida. Use DD/MM/AAAA ou AAAA-MM-DD."]})


def _parse_int(value: Any, line: int, field: str) -> int:
    text = _clean(value)
    try:
        parsed = int(text)
    except ValueError as exc:
        raise DRFValidationError({"rows": [f"Linha {line}: {field} deve ser um número inteiro."]}) from exc
    if parsed <= 0:
        raise DRFValidationError({"rows": [f"Linha {line}: {field} deve ser maior que zero."]})
    return parsed


def _parse_predecessors(value: Any, line: int) -> list[int]:
    text = _clean(value)
    if not text:
        return []
    parts = [part.strip() for part in re.split(r"[;,]", text) if part.strip()]
    predecessor_ids = [_parse_int(part, line, "Predecessores") for part in parts]
    if len(predecessor_ids) != len(set(predecessor_ids)):
        raise DRFValidationError({"rows": [f"Linha {line}: predecessores repetidos."]})
    return predecessor_ids


def _read_csv(uploaded_file) -> list[dict[str, Any]]:
    content = uploaded_file.read().decode("utf-8-sig")
    return list(csv.DictReader(StringIO(content)))


def _read_xlsx(uploaded_file) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DRFValidationError(
            {"file": "Suporte a .xlsx exige instalar a dependência openpyxl."}
        ) from exc

    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    return [dict(zip(headers, row, strict=False)) for row in rows[1:] if any(cell not in [None, ""] for cell in row)]


def _read_uploaded_file(uploaded_file) -> list[dict[str, Any]]:
    extension = Path(uploaded_file.name).suffix.lower()
    if extension == ".csv":
        return _read_csv(uploaded_file)
    if extension == ".xlsx":
        return _read_xlsx(uploaded_file)
    raise DRFValidationError({"file": "Formato não suportado. Envie .xlsx ou .csv."})


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for header, value in row.items():
        field = HEADER_MAP.get(_normalize_header(header))
        if field and field not in normalized:
            normalized[field] = value
    return normalized


def _resolve_user_by_email(project: Project, email: str, role: str, line: int, field: str) -> Usuario:
    user = Usuario.objects.filter(email__iexact=email).first()
    if not user:
        raise DRFValidationError({"rows": [f"Linha {line}: {field} não encontrado: {email}."]})
    if not Membership.objects.filter(usuario=user, projeto=project, papel__codigo=role).exists():
        raise DRFValidationError(
            {"rows": [f"Linha {line}: {field} não possui papel compatível no projeto: {email}."]}
        )
    return user


def _resolve_node(project: Project, module: str, process: str, line: int) -> NoHierarquia:
    try:
        module_node = obter_ou_criar_modulo(project, module)
    except DRFValidationError as exc:
        raise DRFValidationError({"rows": [f"Linha {line}: módulo inválido: {module}."]}) from exc

    if project.modo == Project.Modo.CUTOVER:
        return module_node

    if not process:
        raise DRFValidationError({"rows": [f"Linha {line}: processo é obrigatório para projeto UAT."]})
    try:
        return obter_ou_criar_processo(project, module_node, process)
    except DRFValidationError as exc:
        raise DRFValidationError({"rows": [f"Linha {line}: processo inválido em {module}: {process}."]}) from exc


def _validate_row(project: Project, normalized: dict[str, Any], line: int) -> ImportedActivityRow:
    missing = [field for field in REQUIRED_FIELDS if not _clean(normalized.get(field))]
    if project.modo == Project.Modo.UAT and not _clean(normalized.get("process")):
        missing.append("process")
    if missing:
        raise DRFValidationError({"rows": [f"Linha {line}: campos obrigatórios ausentes: {', '.join(missing)}."]})

    temporary_id = _parse_int(normalized["temporary_id"], line, "Id lista sequencial")
    planned_start = _parse_date(normalized["planned_start"], line, "Data inicio planejado")
    planned_end = _parse_date(normalized["planned_end"], line, "Data final planejada")
    if planned_start > planned_end:
        raise DRFValidationError({"rows": [f"Linha {line}: data inicial não pode ser maior que data final."]})

    predecessor_temporary_ids = _parse_predecessors(normalized.get("predecessors"), line)
    if temporary_id in predecessor_temporary_ids:
        raise DRFValidationError({"rows": [f"Linha {line}: atividade não pode depender dela mesma."]})

    return ImportedActivityRow(
        line=line,
        temporary_id=temporary_id,
        node=_resolve_node(
            project,
            _clean(normalized["module"]),
            _clean(normalized.get("process")),
            line,
        ),
        name=_clean(normalized["name"]),
        tester=_resolve_user_by_email(
            project,
            _clean(normalized["tester_email"]),
            Papel.Codigo.TESTER,
            line,
            "Tester email",
        ),
        developer=_resolve_user_by_email(
            project,
            _clean(normalized["developer_email"]),
            Papel.Codigo.DEV,
            line,
            "Dev email",
        ),
        planned_start=planned_start,
        planned_end=planned_end,
        predecessor_temporary_ids=predecessor_temporary_ids,
        area=_clean(normalized.get("area")),
        system=_clean(normalized.get("system")),
        transaction=_clean(normalized.get("transaction")),
        expected_result=_clean(normalized.get("expected_result")),
        notes=_clean(normalized.get("notes")),
    )


def _validate_import_rows(project: Project, raw_rows: list[dict[str, Any]]) -> list[ImportedActivityRow]:
    if not raw_rows:
        raise DRFValidationError({"file": "Arquivo não possui linhas de atividades."})

    validated_rows: list[ImportedActivityRow] = []
    errors: list[str] = []
    seen_temporary_ids: set[int] = set()

    for index, raw_row in enumerate(raw_rows, start=2):
        try:
            imported_row = _validate_row(project, _normalize_row(raw_row), index)
        except DRFValidationError as exc:
            row_errors = exc.detail.get("rows") if isinstance(exc.detail, dict) else None
            errors.extend(str(error) for error in (row_errors or exc.detail))
            continue

        if imported_row.temporary_id in seen_temporary_ids:
            errors.append(f"Linha {index}: Id lista sequencial repetido: {imported_row.temporary_id}.")
        seen_temporary_ids.add(imported_row.temporary_id)
        validated_rows.append(imported_row)

    known_temporary_ids = {row.temporary_id for row in validated_rows}
    for row in validated_rows:
        missing_predecessors = [
            predecessor_id
            for predecessor_id in row.predecessor_temporary_ids
            if predecessor_id not in known_temporary_ids
        ]
        if missing_predecessors:
            errors.append(
                f"Linha {row.line}: predecessores não encontrados no arquivo: {', '.join(map(str, missing_predecessors))}."
            )

    if errors:
        raise DRFValidationError({"rows": errors})
    return validated_rows


@transaction.atomic
def importar_activities_de_arquivo(*, projeto: Project, uploaded_file) -> list[Activity]:
    rows = _validate_import_rows(projeto, _read_uploaded_file(uploaded_file))
    activities_by_temporary_id: dict[int, Activity] = {}
    created_activities: list[Activity] = []

    for row in rows:
        activity = Activity(
            projeto=projeto,
            no=row.node,
            nome=row.name,
            status=Activity.Status.AGUARDANDO if row.predecessor_temporary_ids else Activity.Status.LIBERADO,
            tester=row.tester,
            desenvolvedor=row.developer,
            data_inicio_planejada=row.planned_start,
            data_conclusao_planejada=row.planned_end,
            area=row.area,
            sistema=row.system,
            transacao=row.transaction,
            resultado_esperado=row.expected_result,
            observacoes=row.notes,
        )
        salvar_activity(activity)
        activities_by_temporary_id[row.temporary_id] = activity
        created_activities.append(activity)

    for row in rows:
        activity = activities_by_temporary_id[row.temporary_id]
        for predecessor_temporary_id in row.predecessor_temporary_ids:
            ActivityPredecessor.objects.create(
                atividade=activity,
                predecessora=activities_by_temporary_id[predecessor_temporary_id],
            )

    return created_activities
